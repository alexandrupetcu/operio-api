"""
templatize-v2 helper: build a DOCX template 1:1 from the Distrigaz FD-691-327
„Tabel de examinare vizuală a sudurilor” spreadsheet (.xls → .xlsx).

The sheet is a form, not data: every visible row becomes one paragraph (or
table row) with the row's exact height, text sits at its column's offset, the
grid is a real Word table with the sheet's column widths (hidden column E
dropped) and borders. Images (Distrigaz logo, letterhead footer) are taken from
the workbook. The sample welds become a `{#suduri}` row loop.

Called by templatize-v2.py for members whose source is a spreadsheet.
"""
from __future__ import annotations

import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from lxml import etree

EMU_PER_TWIP = 635
XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def col_twips(width_chars: float) -> int:
    """Excel column width (characters, max digit width 7 px) → twips."""
    return round((width_chars * 7 + 5) * 15)


def pt(v: float) -> int:
    return round(v * 20)


# ───────────── tiny WordprocessingML writers ─────────────

def rpr(font="Arial", size: float = 12, bold=False, italic=False) -> str:
    s = f'<w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:cs="{font}"/>'
    if bold:
        s += "<w:b/><w:bCs/>"
    if italic:
        s += "<w:i/><w:iCs/>"
    s += f'<w:sz w:val="{round(size * 2)}"/><w:szCs w:val="{round(size * 2)}"/>'
    return f"<w:rPr>{s}</w:rPr>"


def run(text: str, **kw) -> str:
    if text == "\t":
        return f"<w:r>{rpr(**kw)}<w:tab/></w:r>"
    if text == "\n":
        return f"<w:r>{rpr(**kw)}<w:br/></w:r>"
    return f'<w:r>{rpr(**kw)}<w:t xml:space="preserve">{escape(text)}</w:t></w:r>'


def para(runs: str = "", *, height_pt: float | None = None, ind: int = 0, jc: str | None = None, tabs=(),
         exact=True, mark=None, keep_next=False) -> str:
    ppr = ""
    if keep_next:
        ppr += "<w:keepNext/>"
    if tabs:
        ppr += "<w:tabs>" + "".join(f'<w:tab w:val="left" w:pos="{p}"/>' for p in tabs) + "</w:tabs>"
    sp = '<w:spacing w:before="0" w:after="0"'
    if height_pt:
        sp += f' w:line="{pt(height_pt)}" w:lineRule="{"exact" if exact else "atLeast"}"'
    ppr += sp + "/>"
    if ind:
        ppr += f'<w:ind w:left="{ind}"/>'
    if jc:
        ppr += f'<w:jc w:val="{jc}"/>'
    ppr += mark or rpr(size=10)
    return f"<w:p><w:pPr>{ppr}</w:pPr>{runs}</w:p>"


def inline_image(rid: str, cx: int, cy: int, pid: int, name: str) -> str:
    return (
        f'<w:r><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
        f'<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:docPr id="{pid}" name="{name}"/>'
        f'<wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>'
        f'<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        f'<pic:pic><pic:nvPicPr><pic:cNvPr id="{pid}" name="{name}"/><pic:cNvPicPr/></pic:nvPicPr>'
        f'<pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
        f'<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic>'
        f"</a:graphicData></a:graphic></wp:inline></w:drawing></w:r>"
    )


def cell(width: int, content: str, *, span=1, borders="tblr", valign="bottom", shade=None) -> str:
    b = ""
    for side, key in (("top", "t"), ("left", "l"), ("bottom", "b"), ("right", "r")):
        val = "single" if key in borders else "nil"
        b += f'<w:{side} w:val="{val}"' + (' w:sz="4" w:space="0" w:color="000000"/>' if val == "single" else "/>")
    tcpr = f'<w:tcW w:w="{width}" w:type="dxa"/>'
    if span > 1:
        tcpr += f'<w:gridSpan w:val="{span}"/>'
    tcpr += f"<w:tcBorders>{b}</w:tcBorders>"
    tcpr += f'<w:vAlign w:val="{valign}"/>'
    return f"<w:tc><w:tcPr>{tcpr}</w:tcPr>{content}</w:tc>"


def row(height_pt: float, cells: str, exact=True) -> str:
    rule = "exact" if exact else "atLeast"
    return f'<w:tr><w:trPr><w:trHeight w:val="{pt(height_pt)}" w:hRule="{rule}"/><w:cantSplit/></w:trPr>{cells}</w:tr>'


NSDECL = (
    'xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:o="urn:schemas-microsoft-com:office:office" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:wp14="http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:w10="urn:schemas-microsoft-com:office:word" '
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" '
    'xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml" '
    'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
    'mc:Ignorable="w14 w15 wp14"'
)


def sheet_images(xlsx: Path) -> list[tuple[bytes, int, int, int, int]]:
    """[(png, cx, cy, fromCol, fromRow)] for the pictures of sheet 1, in drawing order."""
    out = []
    with zipfile.ZipFile(xlsx) as z:
        drawing = etree.fromstring(z.read("xl/drawings/drawing1.xml"))
        rels = etree.fromstring(z.read("xl/drawings/_rels/drawing1.xml.rels"))
        target = {r.get("Id"): r.get("Target") for r in rels}
        for anchor in drawing:
            pic = anchor.find(f"{{{XDR}}}pic")
            if pic is None:
                continue
            blip = pic.find(f".//{{{A}}}blip")
            ext = pic.find(f".//{{{A}}}xfrm/{{{A}}}ext")
            frm = anchor.find(f"{{{XDR}}}from")
            tgt = target[blip.get(f"{{{RNS}}}embed")]
            path = tgt.lstrip("/") if tgt.startswith("/") else ("xl/drawings/" + tgt).replace("xl/drawings/../", "xl/")
            out.append((z.read(path), int(ext.get("cx")), int(ext.get("cy")),
                        int(frm.find(f"{{{XDR}}}col").text), int(frm.find(f"{{{XDR}}}row").text)))
    return out


def build_xlsx_member(xlsx: Path, out: Path, member: dict, roles: dict, log: list[str]) -> None:
    import openpyxl

    T = member["xlsx"]["texts"]
    wb = openpyxl.load_workbook(xlsx)
    ws = wb.worksheets[0]
    cw = {k: col_twips(v.width) for k, v in ws.column_dimensions.items() if not v.hidden}
    xA, wB, wC, wD, wF = cw["A"], cw["B"], cw["C"], cw["D"], cw["F"]
    xB, xC, xD = xA, xA + wB, xA + wB + wC
    rh = {k: v.height for k, v in ws.row_dimensions.items() if v.height}
    m = ws.page_margins
    imgs = sheet_images(xlsx)
    logo, footer_img = imgs[0], imgs[1]
    log.append(f"xlsx columns B/C/D/F = {wB}/{wC}/{wD}/{wF} twips, {len(imgs)} images")

    body = []
    # rows 1–3: the logo (anchored at B1)
    body.append(para(inline_image("rIdImg1", logo[1], logo[2], 1, "Logo"), height_pt=rh[1] + rh[2] + rh[3], ind=xB))
    # rows 4–8: the agency text box (Arial 12 first line, Calibri 11 bold)
    # The text box starts 2 pt into row 4 and sets its lines tight (Arial 12 + Calibri 11 bold);
    # what is left of rows 4–8 stays as space under it, so the title lands where it was.
    ag = T["agency"]
    box = [13.8] + [11.0] * (len(ag) - 1)
    body.append(para("", height_pt=2))
    body.append(para(run(ag[0], size=12), height_pt=box[0], ind=xB + 50))
    for i, line in enumerate(ag[1:], start=1):
        body.append(para(run(line, font="Calibri", size=11, bold=True), height_pt=box[i], ind=xB + 50))
    body.append(para("", height_pt=sum(rh[r] for r in range(4, 9)) - 2 - sum(box)))
    # row 9: title (C9, leading spaces kept), row 10: Nr/data (D10)
    body.append(para(run(ws["C9"].value, size=14), height_pt=rh[9], ind=xC))
    body.append(para(run(ws["D10"].value, size=12), height_pt=rh[10], ind=xD))
    body.append(para("", height_pt=rh[11]))
    # row 12: Produsul | PE100 SDR11 (bold, C12)
    body.append(para(run("Produsul", size=12) + run("\t", size=12) + run(ws["C12"].value, size=12, bold=True),
                     height_pt=rh[12], ind=xB, tabs=(xC,)))
    # row 13 (B13:F13 merged, wrap, centered vertically): Subansamblul … / poziţie montaj …
    body.append(para(run("Subansamblul ", size=12) + run(T["subansamblu"], size=12, bold=True) + run("\n", size=12)
                     + run(" poziţie  montaj ", size=12) + run(T["pozitie"], size=12, bold=True),
                     height_pt=rh[13] / 2, ind=xB))
    # row 14: Nr.proiect montaj … , nr. schemă izometrică …
    body.append(para(run("Nr.proiect montaj ", size=12) + run(T["proiect"], size=12, bold=True)
                     + run(", nr. schemă izometrică ", size=12) + run(T["schema"], size=12, bold=True) + run(".", size=12),
                     height_pt=rh[14], ind=xB))
    # row 15 is hidden in the sheet; row 16
    body.append(para(run(ws["B16"].value, size=12), height_pt=rh[16], ind=xB))
    body.append(para("", height_pt=rh[17]))

    # rows 18–38: the grid
    grid = f'<w:tblGrid><w:gridCol w:w="{wB}"/><w:gridCol w:w="{wC}"/><w:gridCol w:w="{wD}"/><w:gridCol w:w="{wF}"/></w:tblGrid>'
    tpr = (f'<w:tblPr><w:tblW w:w="{wB + wC + wD + wF}" w:type="dxa"/><w:tblInd w:w="{xB}" w:type="dxa"/>'
           '<w:tblLayout w:type="fixed"/><w:tblCellMar><w:left w:w="20" w:type="dxa"/><w:right w:w="20" w:type="dxa"/></w:tblCellMar></w:tblPr>')
    p12 = lambda t, **k: para(run(t, size=12, **k) if t else "", mark=rpr(size=12))  # noqa: E731
    p14c = lambda t: para(run(t, size=14, bold=True) if t else "", jc="center", mark=rpr(size=14, bold=True))  # noqa: E731
    rows = []
    # header: rows 18, 19, 20 — one frame per column, no inner horizontal rules
    rows.append(row(rh[18], cell(wB, p12("Nr."), borders="tlr") + cell(wC, p12("Număr de ordine "), borders="t")
                    + cell(wD, p12(""), borders="tlr") + cell(wF, p12("Rezultatul  examenului"), borders="tlr")))
    rows.append(row(rh[19], cell(wB, p12("crt."), borders="lr") + cell(wC, p12("al  sudurii"), borders="")
                    + cell(wD, p12("   Defecte constatate"), borders="lr") + cell(wF, p12("    admis / respins"), borders="lr")))
    rows.append(row(rh[20], cell(wB, p12(""), borders="blr") + cell(wC, p12(""), borders="b")
                    + cell(wD, p12(""), borders="blr") + cell(wF, p12(""), borders="blr")))
    # the weld loop row (sample rows 21–23) + the empty rows 24–34 kept
    rows.append(row(rh[21], cell(wB, p14c("{#suduri}{nr_crt}")) + cell(wC, p14c("{cod}")) + cell(wD, p14c("-"))
                    + cell(wF, p14c("{rezultat}{/suduri}"))))
    sample = sum(1 for r in range(21, 35) if ws.cell(r, 3).value not in (None, ""))
    for r in range(21 + sample, 35):
        rows.append(row(rh[r], cell(wB, p14c("")) + cell(wC, p14c("")) + cell(wD, p14c("")) + cell(wF, p14c(""))))
    log.append(f"{sample} sample weld rows → 1 loop row, {34 - 20 - sample} empty rows kept")
    # signatures: rows 35–38 (B:C merged)
    rows.append(row(rh[35], cell(wB + wC, p12("Functia"), span=2) + cell(wD, para(run("Nume prenume", size=12), jc="center", mark=rpr(size=12)))
                    + cell(wF, p12("Semnatura,"))))
    rows.append(row(rh[36], cell(wB + wC, p12("Instalator autorizat"), span=2)
                    + cell(wD, para(run(T["instalator"], size=14, bold=True), mark=rpr(size=14, bold=True))) + cell(wF, p12(""))))
    rows.append(row(rh[37], cell(wB + wC, p12("RTS"), span=2)
                    + cell(wD, para(run(T["rts"], size=12, bold=True), mark=rpr(size=12, bold=True))) + cell(wF, p12(""))))
    rows.append(row(rh[38], cell(wB + wC, p12(""), span=2) + cell(wD, p12("")) + cell(wF, p12(""))))
    body.append(f"<w:tbl>{tpr}{grid}{''.join(rows)}</w:tbl>")
    # Word needs a paragraph after a table; keep it 1 pt so extra weld rows have room
    body.append(para("", height_pt=1, mark=rpr(size=1)))

    # Letterhead footer: the in-sheet code (B40, Arial 10 bold italic), the Distrigaz footer picture
    # (rows 41–45) and the sheet's page footer „FD - 691-327”.
    footer = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:ftr {NSDECL}>'
        + para(run(ws["B40"].value, size=10, bold=True, italic=True), ind=xB)
        + para(inline_image("rIdImg2", footer_img[1], footer_img[2], 3, "Subsol"), ind=xB)
        + para(run("FD - 691-327", size=10))
        + "</w:ftr>"
    )

    def inch(v):
        return round(v * 1440)

    ov = member.get("pageOverride", {})
    mar = {"top": max(inch(m.top), 567), "right": max(inch(m.right), 567), "bottom": max(inch(m.bottom), 567),
           "left": max(inch(m.left), 567), "header": inch(m.header), "footer": inch(m.footer)}
    for k, v in ov.items():
        log.append(f"pgMar {k} {mar[k]} → {v} (override)")
        mar[k] = v
    sect = (
        f'<w:sectPr><w:footerReference w:type="default" r:id="rIdFtr1"/>'
        f'<w:pgSz w:w="11906" w:h="16838" w:code="9"/>'
        f'<w:pgMar w:top="{mar["top"]}" w:right="{mar["right"]}" w:bottom="{mar["bottom"]}" '
        f'w:left="{mar["left"]}" w:header="{mar["header"]}" w:footer="{mar["footer"]}" w:gutter="0"/>'
        f'<w:cols w:space="720"/></w:sectPr>'
    )
    document = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {NSDECL}><w:body>'
                + "".join(body) + sect + "</w:body></w:document>")

    styles = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial" w:eastAsia="Arial"/>'
        '<w:sz w:val="20"/><w:szCs w:val="20"/><w:lang w:val="ro-RO"/></w:rPr></w:rPrDefault>'
        '<w:pPrDefault><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr></w:pPrDefault></w:docDefaults>'
        '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/><w:qFormat/></w:style>'
        '<w:style w:type="table" w:default="1" w:styleId="TableNormal"><w:name w:val="Normal Table"/><w:tblPr>'
        '<w:tblInd w:w="0" w:type="dxa"/><w:tblCellMar><w:top w:w="0" w:type="dxa"/><w:left w:w="108" w:type="dxa"/>'
        '<w:bottom w:w="0" w:type="dxa"/><w:right w:w="108" w:type="dxa"/></w:tblCellMar></w:tblPr></w:style>'
        "</w:styles>"
    )
    settings = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:settings xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:defaultTabStop w:val="720"/><w:characterSpacingControl w:val="doNotCompress"/>'
        '<w:compat><w:compatSetting w:name="compatibilityMode" w:uri="http://schemas.microsoft.com/office/word" w:val="15"/></w:compat>'
        "</w:settings>"
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/><Default Extension="png" ContentType="image/png"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
        '<Override PartName="/word/settings.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml"/>'
        '<Override PartName="/word/footer1.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml"/>'
        '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
        "</Types>"
    )
    root_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>'
        "</Relationships>"
    )
    doc_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdStyles" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
        '<Relationship Id="rIdSettings" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/settings" Target="settings.xml"/>'
        '<Relationship Id="rIdFtr1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer" Target="footer1.xml"/>'
        '<Relationship Id="rIdImg1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image1.png"/>'
        "</Relationships>"
    )
    ftr_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdImg2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image2.png"/>'
        "</Relationships>"
    )
    core = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>FD-691-327 Tabel de examinare vizuală a sudurilor</dc:title></cp:coreProperties>'
    )
    # sanity: every XML part must parse
    for x in (document, footer, styles, settings):
        etree.fromstring(x.encode())
    out.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", root_rels)
        z.writestr("docProps/core.xml", core)
        z.writestr("word/document.xml", document)
        z.writestr("word/styles.xml", styles)
        z.writestr("word/settings.xml", settings)
        z.writestr("word/footer1.xml", footer)
        z.writestr("word/_rels/document.xml.rels", doc_rels)
        z.writestr("word/_rels/footer1.xml.rels", ftr_rels)
        z.writestr("word/media/image1.png", logo[0])
        z.writestr("word/media/image2.png", footer_img[0])
    log.append("built from spreadsheet (A4 portrait, as the sheet's page setup)")
